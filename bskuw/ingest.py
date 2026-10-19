"""Load tracking, SKU attributes, and wholesale sales into DuckDB."""

import csv
import tempfile
import zipfile
from pathlib import Path

from bskuw import db

TRACKING_COLUMNS = (
    ("year_week", "year_week"),
    ("basket", "basket"),
    ("product bs", "product BS"),
    ("sku", "SKU"),
    ("stockqty_w", "StockQty_W"),
)

QUANT_COLUMNS = (
    ("sku", "SKU"),
    ("internal category", "Internal Category"),
    ("product name", "Product Name"),
    ("condition", "Condition"),
    ("brand", "Brand"),
    ("model", "Model"),
    ("size", "Size"),
    ("tread", "Tread"),
    ("majority", "Majority"),
    ("location", "Location"),
    ("sales order (int) date", "Sales Order (int) Date"),
    ("sales order (int) number", "Sales Order (int) Number"),
)

# A tire on this location, with a sales order date and number, is a wholesale sale.
WHOLESALE_LOCATION = "Physical Locations / FRD / Utilisation Scrap"


def load_tracking(path, on_progress=None):
    """Replace tracking rows for every week present in this file."""
    path = _existing_file(path, "Tracking")
    suffix = path.suffix.lower()
    if path.name.startswith("~$"):
        raise ValueError("This is an Excel lock file, not the report. Close the file chooser and pick the .xlsx itself.")
    if suffix == ".csv":
        _tracking_stage_from_csv(path)
    elif suffix == ".xlsx":
        _tracking_stage_from_xlsx(path, on_progress=on_progress)
    else:
        raise ValueError("Tracking file must be .xlsx or .csv.")

    staged = db.fetch_one("SELECT count(*) FROM tracking_stage")[0]
    if not staged:
        raise ValueError(
            "No SKU rows found. Rows are kept only when SKU starts with 'SKU' and week and basket are filled in."
        )

    weeks = [row[0] for row in db.fetchall("SELECT DISTINCT year_week FROM tracking_stage ORDER BY year_week")[1]]
    baskets = db.fetch_one("SELECT count(DISTINCT basket) FROM tracking_stage")[0]

    with db.transaction():
        db.execute(
            """
            DELETE FROM tracking
            WHERE year_week IN (SELECT DISTINCT year_week FROM tracking_stage)
            """
        )
        db.execute("INSERT INTO tracking SELECT * FROM tracking_stage")

    return {"rows": staged, "weeks": weeks, "baskets": baskets, "source_file": str(path)}


def load_quant(path):
    """Replace sku_dim and wholesale sales from one quant CSV.

    The file is scanned once. Wholesale rows are tires whose Location is
    WHOLESALE_LOCATION and whose sales order date and number are both filled in.
    """
    path = _existing_file(path, "Quant")
    if path.stat().st_size == 0:
        raise ValueError(
            f"{path.name} is empty. Point this at the real quant export (CSV, about 2 GB)."
        )
    if path.suffix.lower() != ".csv":
        raise ValueError("Quant report must be a .csv file.")

    header = _header(path)
    mapping = _match_columns(header, QUANT_COLUMNS, path.name)
    _quant_raw(path, mapping)
    _sku_stage_from_quant()
    _wholesale_stage_from_quant(path)

    sku_rows = db.fetch_one("SELECT count(*) FROM sku_stage")[0]
    if not sku_rows:
        raise ValueError("The quant file returned no SKUs. The previous tables were left unchanged.")
    wholesale_rows = db.fetch_one("SELECT count(*) FROM wholesale_stage")[0]
    unparsed_dates = db.fetch_one(
        "SELECT count(*) FROM wholesale_stage WHERE sale_date IS NULL"
    )[0]
    last_sale_date = db.fetch_one(
        """
        SELECT max(parse_sale_date(sale_date_raw))
        FROM quant_raw
        WHERE sale_date_raw IS NOT NULL
        """
    )[0]

    db.ensure_quant_report()
    with db.transaction():
        db.execute("DELETE FROM quant_report")
        db.execute(
            "INSERT INTO quant_report (source_file, last_sale_date) VALUES (?, ?)",
            [str(path), last_sale_date],
        )
        db.execute("DELETE FROM sku_dim")
        db.execute(
            """
            INSERT INTO sku_dim (
                sku, name, brand, model, size, tread, condition,
                internal_category, majority, source
            )
            SELECT
                sku, name, brand, model, size, tread, condition,
                internal_category, majority, source
            FROM sku_stage
            """
        )
        db.execute("DELETE FROM wholesale")
        db.execute(
            """
            INSERT INTO wholesale (sku, sale_date, sale_number, source_file)
            SELECT sku, sale_date, sale_number, source_file
            FROM wholesale_stage
            """
        )
    return {
        "rows": sku_rows,
        "wholesale_rows": wholesale_rows,
        "unparsed_dates": unparsed_dates,
        "last_sale_date": last_sale_date,
        "source": "quant",
    }


def capture_quant_last_date(path):
    """Store the latest sales-order date in the whole quant file.

    This reads every row, not only wholesale sales. The loaded tables stay as they are.
    """
    path = _existing_file(path, "Quant")
    header = _header(path)
    mapping = _match_columns(header, QUANT_COLUMNS, path.name)
    column = _quote(mapping["sales order (int) date"])
    last_sale_date = db.fetch_one(
        """
        SELECT max(parse_sale_date(NULLIF(trim(CAST({column} AS VARCHAR)), '')))
        FROM read_csv(?, header = true, all_varchar = true)
        """.format(column=column),
        [str(path)],
    )[0]
    db.ensure_quant_report()
    with db.transaction():
        db.execute("DELETE FROM quant_report")
        db.execute(
            "INSERT INTO quant_report (source_file, last_sale_date) VALUES (?, ?)",
            [str(path), last_sale_date],
        )
    return last_sale_date


def load_mysql():
    """Replace sku_dim from clear_data.clear_sku. Does not load wholesale sales."""
    import pymysql

    from bskuw.settings import mysql_settings

    settings = mysql_settings()
    connection = pymysql.connect(
        host=settings["host"],
        port=settings["port"],
        user=settings["user"],
        password=settings["password"],
        database=settings["database"],
        charset="utf8mb4",
        cursorclass=pymysql.cursors.SSCursor,
    )
    try:
        cursor = connection.cursor()
        cursor.execute(
            """
            SELECT sku, name, brand, model, size, tread, condition
            FROM clear_data.clear_sku
            """
        )
        db.execute("DROP TABLE IF EXISTS sku_stage")
        db.execute(
            """
            CREATE TEMP TABLE sku_stage (
                sku VARCHAR,
                name VARCHAR,
                brand VARCHAR,
                model VARCHAR,
                size VARCHAR,
                tread VARCHAR,
                condition VARCHAR,
                internal_category VARCHAR,
                majority VARCHAR,
                source VARCHAR
            )
            """
        )
        while True:
            batch = cursor.fetchmany(5000)
            if not batch:
                break
            rows = [_mysql_row(row) for row in batch]
            rows = [row for row in rows if row is not None]
            if rows:
                db.executemany(
                    "INSERT INTO sku_stage VALUES (?, ?, ?, ?, ?, ?, ?, NULL, NULL, 'mysql')",
                    rows,
                )
    finally:
        connection.close()

    db.execute("DROP TABLE IF EXISTS sku_stage_grouped")
    db.execute(
        """
        CREATE TEMP TABLE sku_stage_grouped AS
        SELECT
            sku,
            mode(name) AS name,
            mode(brand) AS brand,
            mode(model) AS model,
            mode(size) AS size,
            mode(tread) AS tread,
            mode(condition) AS condition,
            NULL AS internal_category,
            NULL AS majority,
            'mysql' AS source
        FROM sku_stage
        WHERE sku IS NOT NULL
        GROUP BY sku
        """
    )
    db.execute("DROP TABLE sku_stage")
    db.execute("ALTER TABLE sku_stage_grouped RENAME TO sku_stage")
    return _replace_sku_dim("mysql")


def _quant_raw(path, mapping):
    fields = (
        ("sku", "sku"),
        ("product name", "name"),
        ("brand", "brand"),
        ("model", "model"),
        ("size", "size"),
        ("tread", "tread"),
        ("condition", "condition"),
        ("internal category", "internal_category"),
        ("majority", "majority"),
        ("location", "location"),
        ("sales order (int) date", "sale_date_raw"),
        ("sales order (int) number", "sale_number"),
    )
    select_list = ", ".join(
        "NULLIF(trim(CAST({col} AS VARCHAR)), '') AS {alias}".format(
            col=_quote(mapping[key]),
            alias=alias,
        )
        for key, alias in fields
    )
    db.execute("DROP TABLE IF EXISTS quant_raw")
    db.execute(
        """
        CREATE TEMP TABLE quant_raw AS
        SELECT {select_list}
        FROM read_csv(?, header = true, all_varchar = true)
        """.format(select_list=select_list),
        [str(path)],
    )


def _sku_stage_from_quant():
    db.execute("DROP TABLE IF EXISTS sku_stage")
    db.execute(
        """
        CREATE TEMP TABLE sku_stage AS
        SELECT
            sku,
            mode(name) AS name,
            mode(brand) AS brand,
            mode(model) AS model,
            mode(size) AS size,
            mode(tread) AS tread,
            mode(condition) AS condition,
            mode(internal_category) AS internal_category,
            mode(majority) AS majority,
            'quant' AS source
        FROM quant_raw
        WHERE sku IS NOT NULL
        GROUP BY sku
        """
    )


def _wholesale_stage_from_quant(path):
    db.execute("DROP TABLE IF EXISTS wholesale_stage")
    db.execute(
        """
        CREATE TEMP TABLE wholesale_stage AS
        SELECT
            sku,
            parse_sale_date(sale_date_raw) AS sale_date,
            sale_number,
            ? AS source_file
        FROM quant_raw
        WHERE location = ?
          AND sale_date_raw IS NOT NULL
          AND sale_number IS NOT NULL
        """,
        [str(path), WHOLESALE_LOCATION],
    )


def _replace_sku_dim(source):
    staged = db.fetch_one("SELECT count(*) FROM sku_stage")[0]
    if not staged:
        raise ValueError(f"The {source} source returned no SKUs. The previous SKU table was left unchanged.")
    with db.transaction():
        db.execute("DELETE FROM sku_dim")
        db.execute(
            """
            INSERT INTO sku_dim (
                sku, name, brand, model, size, tread, condition,
                internal_category, majority, source
            )
            SELECT
                sku, name, brand, model, size, tread, condition,
                internal_category, majority, source
            FROM sku_stage
            """
        )
    return {"rows": staged, "source": source}


def _tracking_stage_from_csv(path):
    header = _header(path)
    mapping = _match_columns(header, TRACKING_COLUMNS, path.name)
    db.execute("DROP TABLE IF EXISTS tracking_stage")
    db.execute(
        """
        CREATE TEMP TABLE tracking_stage AS
        SELECT
            NULLIF(trim(CAST({week} AS VARCHAR)), '') AS year_week,
            NULLIF(regexp_replace(trim(CAST({basket} AS VARCHAR)), '\\.0$', ''), '') AS basket,
            NULLIF(trim(CAST({product_bs} AS VARCHAR)), '') AS product_bs,
            NULLIF(trim(CAST({sku} AS VARCHAR)), '') AS sku,
            TRY_CAST(NULLIF(trim(CAST({stock} AS VARCHAR)), '') AS DOUBLE) AS stock_qty_w,
            ? AS source_file
        FROM read_csv(?, header = true, all_varchar = true)
        """.format(
            week=_quote(mapping["year_week"]),
            basket=_quote(mapping["basket"]),
            product_bs=_quote(mapping["product bs"]),
            sku=_quote(mapping["sku"]),
            stock=_quote(mapping["stockqty_w"]),
        ),
        [str(path), str(path)],
    )
    _keep_tracking_sku_rows()


def _tracking_stage_from_xlsx(path, on_progress=None):
    if _try_tracking_stage_from_duckdb(path, path):
        _keep_tracking_sku_rows()
        return
    repaired = _plain_styles_copy(path)
    try:
        if repaired is not None and _try_tracking_stage_from_duckdb(repaired, path):
            _keep_tracking_sku_rows()
            return
    finally:
        if repaired is not None:
            try:
                repaired.unlink(missing_ok=True)
            except OSError:
                pass
    _tracking_stage_from_openpyxl(path, on_progress=on_progress)


def _try_tracking_stage_from_duckdb(read_path, source_file):
    """Project the five tracking columns with read_xlsx.

    all_varchar keeps a trailing Total row from aborting the scan. Callers
    drop rows whose SKU does not start with SKU.
    """
    try:
        _load_excel_extension()
        db.execute("DROP TABLE IF EXISTS tracking_stage")
        db.execute(
            """
            CREATE TEMP TABLE tracking_stage AS
            SELECT
                NULLIF(trim(CAST("year_week" AS VARCHAR)), '') AS year_week,
                NULLIF(regexp_replace(trim(CAST("basket" AS VARCHAR)), '\\.0$', ''), '') AS basket,
                NULLIF(trim(CAST("product BS" AS VARCHAR)), '') AS product_bs,
                NULLIF(trim(CAST("SKU" AS VARCHAR)), '') AS sku,
                TRY_CAST("StockQty_W" AS DOUBLE) AS stock_qty_w,
                ? AS source_file
            FROM read_xlsx(?, header = true, all_varchar = true, sheet = 'Export')
            """,
            [str(source_file), str(read_path)],
        )
        return True
    except Exception:
        _clear_failed_read()
        return False


_PLAIN_STYLES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>
<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>
<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
<cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>
<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
</styleSheet>
"""


def _plain_styles_copy(path):
    """Copy an xlsx with a standards-compliant styles part.

    The tracking export's styles.xml has no cellStyleXfs, and DuckDB refuses
    to open it. Values are unchanged; only the style table is replaced.
    """
    temp = Path(tempfile.gettempdir()) / "bskuw-tracking-{pid}.xlsx".format(pid=path.stat().st_mtime_ns)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED) as dest:
            for item in source.infolist():
                payload = _PLAIN_STYLES.encode("utf-8") if item.filename == "xl/styles.xml" else source.read(item.filename)
                dest.writestr(item, payload)
    except Exception:
        temp.unlink(missing_ok=True)
        return None
    return temp


def _clear_failed_read():
    """A failed read_xlsx can leave DuckDB's transaction aborted."""
    try:
        db.execute("ROLLBACK")
    except Exception:
        pass
    try:
        db.execute("DROP TABLE IF EXISTS tracking_stage")
    except Exception:
        pass


def _load_excel_extension():
    connection = db.get_connection()
    try:
        connection.execute("LOAD excel")
    except Exception:
        connection.execute("INSTALL excel")
        connection.execute("LOAD excel")


def _tracking_stage_from_openpyxl(path, on_progress=None):
    from openpyxl import load_workbook

    db.execute("DROP TABLE IF EXISTS tracking_stage")
    db.execute(
        """
        CREATE TEMP TABLE tracking_stage (
            year_week VARCHAR,
            basket VARCHAR,
            product_bs VARCHAR,
            sku VARCHAR,
            stock_qty_w DOUBLE,
            source_file VARCHAR
        )
        """
    )
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Export"] if "Export" in workbook.sheetnames else workbook.active
        row_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(row_iter)
        except StopIteration:
            raise ValueError(f"{path.name} has no header row.")
        names = ["" if cell is None else str(cell).strip() for cell in header]
        mapping = _match_columns(names, TRACKING_COLUMNS, path.name)
        indexes = {key: names.index(mapping[key]) for key in mapping}
        source = str(path)
        batch = []
        kept = 0
        for raw in row_iter:
            record = _tracking_record(raw, indexes, source)
            if record is None:
                continue
            batch.append(record)
            kept += 1
            if len(batch) >= 5000:
                db.executemany("INSERT INTO tracking_stage VALUES (?, ?, ?, ?, ?, ?)", batch)
                batch = []
                if on_progress:
                    on_progress(kept)
        if batch:
            db.executemany("INSERT INTO tracking_stage VALUES (?, ?, ?, ?, ?, ?)", batch)
            if on_progress:
                on_progress(kept)
    finally:
        workbook.close()


def _tracking_record(raw, indexes, source):
    def cell(key):
        position = indexes[key]
        if position >= len(raw):
            return None
        return raw[position]

    sku = _as_text(cell("sku"))
    year_week = _as_text(cell("year_week"))
    basket = _as_text(cell("basket"))
    if not sku or not sku.upper().startswith("SKU") or not year_week or not basket:
        return None
    return (
        year_week,
        basket,
        _as_text(cell("product bs")),
        sku,
        _as_float(cell("stockqty_w")),
        source,
    )


def _keep_tracking_sku_rows():
    db.execute(
        """
        DELETE FROM tracking_stage
        WHERE sku IS NULL
           OR upper(sku) NOT LIKE 'SKU%'
           OR year_week IS NULL
           OR basket IS NULL
        """
    )


def _mysql_row(row):
    sku = _as_text(row[0])
    if not sku:
        return None
    return (
        sku,
        _as_text(row[1]),
        _as_text(row[2]),
        _as_text(row[3]),
        _as_text(row[4]),
        _as_text(row[5]),
        _as_text(row[6]),
    )


def _existing_file(path, label):
    if path is None or not str(path).strip():
        raise ValueError(f"Choose a {label.lower()} file.")
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"{label} file not found: {file_path}")
    return file_path


def _header(path):
    blob = path.open("rb").read(65536)
    if not blob.strip():
        raise ValueError(f"{path.name} is empty.")
    sample = blob.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(sample[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    line = sample.splitlines()[0]
    header = [name.strip() for name in next(csv.reader([line], dialect))]
    if not any(header):
        raise ValueError(f"{path.name} has no header row.")
    return header


def _match_columns(header, required, file_name):
    by_key = {}
    for name in header:
        key = " ".join(name.strip().lower().split())
        if key and key not in by_key:
            by_key[key] = name.strip()
    missing = [label for key, label in required if key not in by_key]
    if missing:
        found = ", ".join(name for name in header if name) or "(none)"
        raise ValueError(
            "{file} is missing columns: {missing}. Found: {found}".format(
                file=file_name,
                missing=", ".join(missing),
                found=found,
            )
        )
    return {key: by_key[key] for key, _label in required}


def _quote(identifier):
    return '"' + identifier.replace('"', '""') + '"'


def _as_text(value):
    if value is None:
        return None
    if isinstance(value, float):
        if value != value:
            return None
        if value.is_integer():
            return str(int(value))
        text = str(value).strip()
        return text or None
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    return text or None


def _as_float(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value != value:
            return None
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    return float(text)
